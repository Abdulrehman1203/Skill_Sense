"""Regression coverage for review findings R11–R15."""

import io
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

from celery.exceptions import Retry
from django.test import SimpleTestCase, override_settings
from rest_framework.test import APIClient, APITestCase

from ai.matching.sbert import match
from interview_system.models import (
    Application, CandidateProfile, ClerkWebhookState, Job, ParsedResume,
    RecruiterProfile, Resume, User,
)
from interview_system.resumes.extraction import ExtractionError, _extract_docx_text, _extract_pdf_text, extract_text
from interview_system.tasks.parsing import compute_match_score
from interview_system.webhook_views import ClerkWebhookView


class ExtractionAndMatchingTests(SimpleTestCase):
    def test_docx_table_nested_table_header_and_footer_are_extracted(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Lead engineer")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Python in table"
        nested = table.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Django in nested table"
        document.sections[0].header.paragraphs[0].text = "Header contact"
        document.sections[0].footer.paragraphs[0].text = "Footer contact"
        stream = io.BytesIO()
        document.save(stream)
        extracted = _extract_docx_text(stream.getvalue())
        for text in ("Lead engineer", "Python in table", "Django in nested table", "Header contact", "Footer contact"):
            self.assertIn(text, extracted)

    def test_library_and_mime_failures_use_extraction_error(self):
        with patch.dict("sys.modules", {"pdfplumber": None}):
            with self.assertRaises(ExtractionError):
                _extract_pdf_text(b"%PDF-1.4")
        with patch("interview_system.resumes.extraction._detect_mime", side_effect=ImportError("magic missing")):
            with self.assertRaises(ExtractionError):
                extract_text(io.BytesIO(b"some bytes"))
        with patch("interview_system.resumes.extraction._detect_mime", return_value="application/pdf"), patch(
            "interview_system.resumes.extraction._extract_pdf_text", side_effect=ImportError("pdf library missing")
        ):
            with self.assertRaises(ExtractionError):
                extract_text(io.BytesIO(b"%PDF-1.4"))

    def test_skill_terms_require_boundaries_and_aliases(self):
        with patch("ai.matching.sbert._embedding_for", return_value=[1.0, 0.0]):
            result = match("JavaScript developer using JS, PostgreSQL and C++", "Backend role", [
                "Java", "R", "JavaScript", "Postgres", "C++", "Python",
            ])
        self.assertEqual(result["matched_skills"], ["JavaScript", "Postgres", "C++"])
        self.assertEqual(result["missing_skills"], ["Java", "R", "Python"])


class ProcessingFailureTests(APITestCase):
    def setUp(self):
        recruiter = User.objects.create(clerk_id="r11_recruiter", email="r11-recruiter@example.test", role="RECRUITER")
        candidate = User.objects.create(clerk_id="r11_candidate", email="r11-candidate@example.test", role="CANDIDATE")
        recruiter_profile = RecruiterProfile.objects.create(user=recruiter)
        candidate_profile = CandidateProfile.objects.create(user=candidate)
        job = Job.objects.create(
            recruiter=recruiter_profile, title="Engineer", description="Python",
            requirements="Python", skills_required=["Python"], location="Remote",
            job_type="REMOTE", experience_level="ENTRY", status=Job.Status.ACTIVE,
        )
        self.resume = Resume.objects.create(candidate=candidate_profile, file="resumes/test.pdf", status=Resume.Status.PARSED)
        Application.objects.create(candidate=candidate_profile, job=job, resume=self.resume)
        self.parsed = ParsedResume.objects.create(
            resume=self.resume, raw_text="Python engineer", match_score=87,
            matched_skills=["Python"], missing_skills=[],
        )
        self.client.force_authenticate(candidate)

    def test_matching_retries_then_records_terminal_failure(self):
        with patch("ai.matching.sbert.match", side_effect=RuntimeError("model unavailable")) as matcher:
            with self.assertRaises(Retry):
                compute_match_score.apply(kwargs={"resume_id": str(self.resume.pk)}, throw=True)
            self.resume.refresh_from_db()
            self.assertEqual(self.resume.status, Resume.Status.PARSED)
            with self.assertRaises(RuntimeError):
                compute_match_score.apply(kwargs={"resume_id": str(self.resume.pk)}, retries=3, throw=True)
        self.assertEqual(matcher.call_count, 2)
        self.resume.refresh_from_db()
        self.parsed.refresh_from_db()
        self.assertEqual(self.resume.status, Resume.Status.FAILED)
        self.assertIn("model unavailable", self.resume.processing_error)
        self.assertIsNone(self.parsed.match_score)
        self.assertEqual(self.parsed.matched_skills, [])
        detail = self.client.get(f"/api/resumes/{self.resume.pk}/")
        self.assertEqual(detail.data["processing_error"], self.resume.processing_error)

    def test_nontransient_matching_error_fails_without_retry(self):
        with patch("ai.matching.sbert.match", side_effect=ValueError("invalid embedding")) as matcher:
            with self.assertRaises(ValueError):
                compute_match_score.apply(kwargs={"resume_id": str(self.resume.pk)}, throw=True)
        self.assertEqual(matcher.call_count, 1)
        self.resume.refresh_from_db()
        self.assertEqual(self.resume.status, Resume.Status.FAILED)


class WebhookOrderingAndContractTests(APITestCase):
    def setUp(self):
        self.view = ClerkWebhookView()
        self.event_at = datetime(2026, 9, 17, 12, 0, tzinfo=timezone.utc)
        self.payload = {
            "id": "r14_user", "email_addresses": [{"id": "email", "email_address": "r14@example.test"}],
            "primary_email_address_id": "email", "first_name": "First",
        }

    def test_delete_then_older_create_does_not_reactivate(self):
        self.view._process_user_event("user.created", self.payload, self.event_at)
        self.view._process_user_event("user.deleted", self.payload, self.event_at + timedelta(seconds=2))
        self.view._process_user_event("user.created", self.payload, self.event_at + timedelta(seconds=1))
        self.view._process_user_event("user.updated", self.payload, self.event_at + timedelta(seconds=3))
        user = User.objects.get(clerk_id="r14_user")
        self.assertFalse(user.is_active)
        self.assertTrue(ClerkWebhookState.objects.get(clerk_id="r14_user").is_deleted)

    def test_delete_before_create_leaves_tombstone(self):
        self.view._process_user_event("user.deleted", self.payload, self.event_at)
        self.view._process_user_event("user.created", self.payload, self.event_at - timedelta(seconds=1))
        self.assertFalse(User.objects.filter(clerk_id="r14_user").exists())

    def test_local_admin_deactivation_survives_replayed_create(self):
        self.view._process_user_event("user.created", self.payload, self.event_at)
        User.objects.filter(clerk_id="r14_user").update(is_active=False)
        self.view._process_user_event("user.created", self.payload, self.event_at + timedelta(seconds=1))
        self.assertFalse(User.objects.get(clerk_id="r14_user").is_active)

    def test_me_requires_bearer_and_returns_401(self):
        response = APIClient().get("/api/users/me/")
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response["WWW-Authenticate"], "Bearer")

    def test_phone_format_applies_to_both_profiles(self):
        recruiter = User.objects.create(clerk_id="phone_recruiter", email="phone-r@example.test", role="RECRUITER")
        candidate = User.objects.create(clerk_id="phone_candidate", email="phone-c@example.test", role="CANDIDATE")
        RecruiterProfile.objects.create(user=recruiter)
        CandidateProfile.objects.create(user=candidate)
        for user, endpoint in ((recruiter, "/api/recruiters/profile/"), (candidate, "/api/candidates/profile/")):
            self.client.force_authenticate(user)
            self.assertEqual(self.client.patch(endpoint, {"phone": "abc123"}, format="json").status_code, 400)
            valid = self.client.patch(endpoint, {"phone": "+923001234567"}, format="json")
            self.assertEqual(valid.status_code, 200)
            self.assertEqual(valid.data["phone"], "+923001234567")
