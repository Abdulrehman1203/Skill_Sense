"""Regression coverage for job-skill access and resume upload lifecycle."""

from __future__ import annotations

import io
import tempfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APIClient, APITestCase

from interview_system.models import (
    Application, CandidateProfile, Job, JobSkill, ParsedResume,
    RecruiterProfile, Resume, User,
)
from interview_system.resumes.extraction import extract_text


def valid_pdf_bytes() -> bytes:
    """A small PDF with an actual extractable text layer."""
    content = (
        b"BT /F1 12 Tf 72 720 Td "
        b"(Python developer with more than ten years of experience building reliable systems) "
        b"Tj ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(offsets)}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010d} 00000 n \n".encode()
    pdf += (
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return pdf


class CurrentPhaseRegressionTests(APITestCase):
    def setUp(self):
        media = tempfile.TemporaryDirectory(prefix="skillsense-test-media-")
        self.addCleanup(media.cleanup)
        settings_override = override_settings(MEDIA_ROOT=media.name)
        settings_override.enable()
        self.addCleanup(settings_override.disable)

        self.owner = User.objects.create(clerk_id="fix_owner", email="owner@example.test", role="RECRUITER")
        self.other = User.objects.create(clerk_id="fix_other", email="other@example.test", role="RECRUITER")
        self.candidate = User.objects.create(clerk_id="fix_candidate", email="candidate@example.test")
        owner_profile = RecruiterProfile.objects.create(user=self.owner)
        RecruiterProfile.objects.create(user=self.other)
        self.candidate_profile = CandidateProfile.objects.create(user=self.candidate)
        self.draft = Job.objects.create(
            recruiter=owner_profile, title="Draft", description="Python", requirements="Python",
            skills_required=["Python"], location="Remote", job_type="REMOTE",
            experience_level="ENTRY", status=Job.Status.DRAFT,
        )
        self.active = Job.objects.create(
            recruiter=owner_profile, title="Active", description="Python", requirements="Python",
            skills_required=["Python"], location="Remote", job_type="REMOTE",
            experience_level="ENTRY", status=Job.Status.ACTIVE,
        )
        self.client = APIClient()

    def test_other_recruiter_cannot_mutate_skills(self):
        skill = JobSkill.objects.create(job=self.draft, skill_name="Python")
        base = f"/api/jobs/{self.draft.pk}/skills/"
        detail = f"{base}{skill.pk}/"
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post(base, {"skill_name": "Rust"}).status_code, 404)
        self.assertEqual(self.client.patch(detail, {"skill_name": "Rust"}).status_code, 404)
        self.assertEqual(self.client.delete(detail).status_code, 404)
        self.assertFalse(JobSkill.objects.filter(job=self.draft, skill_name="Rust").exists())
        self.assertTrue(JobSkill.objects.filter(pk=skill.pk).exists())

    def test_draft_skills_visible_only_to_owner_and_active_skills_public(self):
        JobSkill.objects.create(job=self.draft, skill_name="PrivateSkill")
        JobSkill.objects.create(job=self.active, skill_name="PublicSkill")
        draft_url = f"/api/jobs/{self.draft.pk}/skills/"
        active_url = f"/api/jobs/{self.active.pk}/skills/"
        self.assertEqual(self.client.get(draft_url).status_code, 404)
        self.assertEqual(self.client.get(active_url).status_code, 200)
        self.client.force_authenticate(self.candidate)
        self.assertEqual(self.client.get(draft_url).status_code, 404)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(draft_url).status_code, 404)
        self.client.force_authenticate(self.owner)
        self.assertEqual(self.client.get(draft_url).status_code, 200)
        self.assertEqual(self.client.post(draft_url, {"skill_name": "Django"}).status_code, 201)

    def test_submitted_resume_cannot_be_replaced(self):
        resume = Resume.objects.create(
            candidate=self.candidate_profile,
            file=SimpleUploadedFile("old.pdf", valid_pdf_bytes()),
            status=Resume.Status.PARSED,
        )
        ParsedResume.objects.create(resume=resume, raw_text="Old Python text", match_score=50)
        Application.objects.create(candidate=self.candidate_profile, job=self.active, resume=resume)
        old_name = resume.file.name
        self.client.force_authenticate(self.candidate)
        response = self.client.patch(
            f"/api/candidates/resumes/{resume.pk}/",
            {"file": SimpleUploadedFile("new.pdf", valid_pdf_bytes())},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("file", response.data)
        resume.refresh_from_db()
        self.assertEqual(resume.file.name, old_name)
        self.assertEqual(resume.parsed_data.raw_text, "Old Python text")

    def test_standalone_uploads_reject_bad_size_extension_and_content(self):
        self.client.force_authenticate(self.candidate)
        invalid = (
            ("big.pdf", b"%PDF-" + b"x" * (5 * 1024 * 1024)),
            ("legacy.doc", b"not supported"),
            ("fake.pdf", b"not actually a pdf"),
            ("fake.docx", b"PK\x03\x04not really an archive"),
        )
        for name, data in invalid:
            with self.subTest(name=name):
                response = self.client.post(
                    "/api/candidates/resumes/",
                    {"file": SimpleUploadedFile(name, data)},
                    format="multipart",
                )
                self.assertEqual(response.status_code, 400)
        self.assertEqual(Resume.objects.count(), 0)

    def test_standalone_upload_is_stored_without_scheduling_processing(self):
        self.client.force_authenticate(self.candidate)
        with self.captureOnCommitCallbacks(execute=True) as callbacks, patch(
            "interview_system.signals.parse_resume.delay"
        ) as parse_task:
            response = self.client.post(
                "/api/candidates/resumes/",
                {"file": SimpleUploadedFile("stored.pdf", valid_pdf_bytes())},
                format="multipart",
            )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["status"], Resume.Status.STORED)
        self.assertEqual(callbacks, [])
        parse_task.assert_not_called()

    def test_unsubmitted_resume_can_be_replaced_with_valid_docx(self):
        from docx import Document

        self.client.force_authenticate(self.candidate)
        created = self.client.post(
            "/api/candidates/resumes/",
            {"file": SimpleUploadedFile("original.pdf", valid_pdf_bytes())},
            format="multipart",
        )
        self.assertEqual(created.status_code, 201)
        document = Document()
        document.add_paragraph("Python developer")
        stream = io.BytesIO()
        document.save(stream)
        response = self.client.patch(
            f"/api/candidates/resumes/{created.data['id']}/",
            {"file": SimpleUploadedFile("replacement.docx", stream.getvalue())},
            format="multipart",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["status"], Resume.Status.STORED)
        resume = Resume.objects.get(pk=created.data["id"])
        self.assertTrue(resume.file.name.endswith(".docx"))

    def test_application_upload_remains_pending_and_queues_parsing(self):
        self.client.force_authenticate(self.candidate)
        with patch("interview_system.signals.parse_resume.delay") as parse_task, \
                self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                "/api/applications/",
                {
                    "job": str(self.active.pk),
                    "resume_file": SimpleUploadedFile("submitted.pdf", valid_pdf_bytes()),
                    "consent_given": True,
                },
                format="multipart",
            )
        self.assertEqual(response.status_code, 201)
        application = Application.objects.get(pk=response.data["id"])
        self.assertEqual(application.resume.status, Resume.Status.PENDING)
        parse_task.assert_called_once_with(resume_id=str(application.resume_id))

    def test_real_pdf_and_docx_extract_text(self):
        from docx import Document

        self.assertIn("Python developer", extract_text(io.BytesIO(valid_pdf_bytes())))
        document = Document()
        document.add_paragraph("Python developer")
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        self.assertIn("Python developer", extract_text(stream))

    def test_storage_file_handle_is_released_after_extraction(self):
        resume = Resume.objects.create(
            candidate=self.candidate_profile,
            file=SimpleUploadedFile("stored.pdf", valid_pdf_bytes()),
            status=Resume.Status.STORED,
        )
        self.assertIn("Python developer", extract_text(resume.file))
        self.assertTrue(resume.file.closed)
